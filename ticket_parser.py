"""Conservative extraction for one labelled concrete delivery ticket per file."""
import argparse
from datetime import date
import hashlib
import html
import json
import math
from pathlib import Path
import re
import sqlite3

ALIASES = {
    'ticket no': 'ticket_id', 'ticket number': 'ticket_id',
    'delivery date': 'delivery_date', 'date': 'delivery_date',
    'supplier': 'supplier', 'site': 'site', 'mix': 'mix', 'mix reference': 'mix',
    'quantity': 'quantity', 'volume': 'quantity',
}
REQUIRED = ('ticket_id', 'delivery_date', 'supplier', 'site', 'mix', 'quantity')


def parse_text(text, source='pasted-text'):
    fields, issues, evidence = {}, [], {}
    for line_no, line in enumerate(text.splitlines(), 1):
        match = re.fullmatch(r'\s*([^:]+):\s*(.*?)\s*', line)
        if not match:
            continue
        key = ALIASES.get(match[1].strip().lower())
        if key:
            if key in fields:
                issues.append('Repeated field: ' + key)
            else:
                fields[key] = match[2]
                evidence[key] = {'line': line_no, 'text': line.strip()}
    for key in REQUIRED:
        if not fields.get(key):
            issues.append('Missing field: ' + key)
    if fields.get('delivery_date'):
        try:
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', fields['delivery_date']):
                raise ValueError()
            date.fromisoformat(fields['delivery_date'])
        except ValueError:
            issues.append('Date must be a valid YYYY-MM-DD value')
    if fields.get('ticket_id') and not re.fullmatch(r'[A-Za-z0-9_-]{1,60}', fields['ticket_id']):
        issues.append('Ticket ID must contain only letters, digits, dash or underscore')
    quantity = re.fullmatch(r'(\d+(?:\.\d+)?)\s*(m3|m³|litres?|liters?|l)',
                            fields.get('quantity', ''), re.IGNORECASE)
    fields['volume_m3'] = None
    if quantity:
        volume = float(quantity[1]) / (1 if quantity[2].lower() in ('m3', 'm³') else 1000)
        if math.isfinite(volume) and volume > 0:
            fields['volume_m3'] = volume
        else:
            issues.append('Volume must be finite and positive')
    elif fields.get('quantity'):
        issues.append('Quantity requires a positive number with m3, m³ or litre units')
    return {'source': source, 'sha256': hashlib.sha256(text.encode()).hexdigest(),
            'status': 'review' if issues else 'validated_for_export',
            'fields': fields, 'issues': issues, 'evidence': evidence}


def read_ticket(path):
    path = Path(path)
    try:
        if path.suffix.lower() == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError('Encrypted PDF is unsupported')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
            if not text.strip():
                raise ValueError('No text layer: OCR and manual verification required')
        elif path.suffix.lower() == '.txt':
            text = path.read_text(encoding='utf-8')
        else:
            raise ValueError('Only .txt and text-layer .pdf inputs are supported')
        result = parse_text(text, path.name)
    except Exception as exc:
        result = {'source': path.name, 'status': 'review', 'fields': {},
                  'issues': ['Read error: ' + str(exc)], 'evidence': {}}
    result['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    return result


def process(paths):
    records = [read_ticket(path) for path in sorted(map(Path, paths))]
    groups = {}
    for rec in records:
        f = rec['fields']
        if f.get('supplier') and f.get('ticket_id'):
            key = (f['supplier'].strip().casefold(), f['ticket_id'].casefold())
            groups.setdefault(key, []).append(rec)
    for group in groups.values():
        if len(group) > 1:
            for rec in group:
                rec['status'] = 'review'
                rec['issues'].append('Duplicate supplier/ticket ID in batch; all copies quarantined')
    return records


def export(records, out):
    """Create a fresh batch snapshot. Never append silently to a previous export."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    out = Path(out)
    out.mkdir(parents=True, exist_ok=False)
    columns = ['supplier', 'ticket_id', 'delivery_date', 'site', 'mix', 'volume_m3', 'source', 'sha256']
    accepted = [r for r in records if r['status'] == 'validated_for_export']
    with sqlite3.connect(out / 'tickets.sqlite') as db:
        db.execute('CREATE TABLE tickets (supplier TEXT, ticket_id TEXT, delivery_date TEXT, '
                   'site TEXT, mix TEXT, volume_m3 REAL CHECK(volume_m3>0), source TEXT, '
                   'sha256 TEXT, PRIMARY KEY(supplier,ticket_id))')
        for rec in accepted:
            row = dict(rec['fields'], source=rec['source'], sha256=rec['sha256'])
            db.execute('INSERT INTO tickets VALUES (?,?,?,?,?,?,?,?)', [row[k] for k in columns])
    wb = Workbook()
    ws = wb.active
    ws.title = 'Validated tickets'
    ws.append(columns)
    for rec in accepted:
        row = dict(rec['fields'], source=rec['source'], sha256=rec['sha256'])
        ws.append([row[k] for k in columns])
    review = wb.create_sheet('Review queue')
    review.append(['source', 'issues', 'sha256'])
    for rec in records:
        if rec['status'] == 'review':
            review.append([rec['source'], '; '.join(rec['issues']), rec['sha256']])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='153747')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Preserve =,+,-,@ as text, never as formulas.
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(65, max(18, max(len(str(c.value or '')) for c in col) + 2))
    for cell in ws['F'][1:]:
        cell.number_format = '0.000'
    wb.save(out / 'tickets.xlsx')
    (out / 'audit.json').write_text(json.dumps(records, indent=2, ensure_ascii=False) + '\n')
    total = sum(r['fields']['volume_m3'] for r in accepted)
    summary = f'{len(accepted)} validated; {len(records)-len(accepted)} for review; {total:.3f} m³ validated volume.'
    rows = ''.join('<tr><td>' + html.escape(r['source']) + '</td><td>' + html.escape(r['status']) +
                   '</td><td>' + html.escape('; '.join(r['issues'])) + '</td></tr>' for r in records)
    (out / 'report.html').write_text('<!doctype html><meta charset="utf-8"><title>Ticket review</title>'
        '<style>body{font:16px system-ui;max-width:1000px;margin:40px auto;padding:20px;color:#153747}'
        'td,th{padding:12px;border-bottom:1px solid #ccc;text-align:left}</style>'
        '<h1>Concrete delivery ticket review</h1><p>Synthetic portfolio demo. Human review required.</p><p>' +
        summary + '</p><table><tr><th>Source</th><th>Status</th><th>Review reason</th></tr>' + rows + '</table>')
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('inputs', nargs='+', type=Path)
    parser.add_argument('--out', type=Path, required=True, help='New output directory (must not exist)')
    args = parser.parse_args()
    try:
        print(export(process(args.inputs), args.out))
    except (ValueError, OSError) as exc:
        parser.exit(2, f'Error: {exc}\n')


if __name__ == '__main__':
    main()

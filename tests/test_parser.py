import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfWriter
from make_demo_pdf import make_pdf
from ticket_parser import parse_text, process, read_ticket, export

TEXT = Path('examples/ticket-001.txt').read_text()

class ParserTests(unittest.TestCase):
    def test_valid_ticket(self):
        r = parse_text(TEXT)
        self.assertEqual(r['status'], 'validated_for_export')
        self.assertEqual(r['fields']['volume_m3'], 6.5)
        self.assertEqual(r['evidence']['ticket_id']['line'], 2)
    def test_litre_conversion(self):
        self.assertEqual(parse_text(TEXT.replace('6.5 m3', '6500 litres'))['fields']['volume_m3'], 6.5)
    def test_missing_field(self):
        self.assertIn('Missing field: mix', parse_text(TEXT.replace('Mix: DEMO-C32', ''))['issues'])
    def test_ambiguous_date(self):
        self.assertEqual(parse_text(TEXT.replace('2026-01-12','12/01/26'))['status'], 'review')
    def test_invalid_calendar_date(self):
        self.assertEqual(parse_text(TEXT.replace('2026-01-12','2026-02-30'))['status'], 'review')
    def test_bad_quantities(self):
        for q in ['8', 'nan m3', '-2 m3', '0 m3', '6,5 m3', '6.5 tonnes']:
            with self.subTest(q=q):
                self.assertEqual(parse_text(TEXT.replace('6.5 m3', q))['status'], 'review')
    def test_repeated_field(self):
        self.assertEqual(parse_text(TEXT + 'Ticket No: DEMO-002\n')['status'], 'review')
    def test_duplicate_quarantine(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = [Path(tmp)/'one.txt', Path(tmp)/'two.txt']
            for p in paths: p.write_text(TEXT)
            self.assertTrue(all(r['status']=='review' for r in process(paths)))
    def test_pdf_text_extraction(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)/'ticket.pdf'
            make_pdf(TEXT, p)
            self.assertEqual(read_ticket(p)['fields']['volume_m3'], 6.5)
    def test_pdf_without_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)/'blank.pdf'
            writer = PdfWriter(); writer.add_blank_page(595,842); writer.write(p)
            self.assertIn('OCR', read_ticket(p)['issues'][0])
    def test_export_consistency_and_formula_safety(self):
        records = [parse_text(TEXT.replace('Training slab A', '=1+1')), parse_text('invalid')]
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)/'result'
            export(records, out)
            with sqlite3.connect(out/'tickets.sqlite') as db:
                self.assertEqual(db.execute('SELECT count(*), sum(volume_m3) FROM tickets').fetchone(), (1,6.5))
            wb = load_workbook(out/'tickets.xlsx')
            self.assertEqual(wb.worksheets[0]['D2'].data_type, 's')
            self.assertEqual(wb.worksheets[0]['D2'].value, '=1+1')
            self.assertEqual(wb.worksheets[1].max_row, 2)
            wb.close()
            self.assertEqual(len(json.loads((out/'audit.json').read_text())), 2)
    def test_existing_output_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(FileExistsError): export([], tmp)

if __name__ == '__main__': unittest.main()
